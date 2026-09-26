"""Case Study 4 (Kestrel Industrial Components) — illustrative 5-year TCO and value model.

Builds TCO-Analysis.xlsx: every hardcoded number is a blue input on the Assumptions or
OneTime sheet with its basis stated; every other cell is a formula. All prices are
ILLUSTRATIVE planning assumptions, not vendor quotes. Several (GDC, Manufacturing Connect,
Azure IoT Operations) are quote-based and must be replaced with real quotes.
Year 1 = Oct 2026 – Sep 2027 (M1–M12), matching docs/migration-roadmap.md.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter as L

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
BOLD_WHITE = Font(name=FONT, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, bold=True, size=14)
SECTION = Font(name=FONT, bold=True, size=11)
NOTE = Font(name=FONT, italic=True, color="595959", size=9)
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEADER = PatternFill("solid", fgColor="1F4E79")
SUB = PatternFill("solid", fgColor="D9E1F2")
TOTAL = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="B7B7B7")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CUR = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
NUM = '#,##0;(#,##0);"-"'
DEC = '#,##0.00;(#,##0.00);"-"'
YEARS = ["Y1", "Y2", "Y3", "Y4", "Y5"]

wb = openpyxl.Workbook()


def hdr(ws, row, labels, start=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=start + i, value=t)
        c.font, c.fill, c.border = BOLD_WHITE, HEADER, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[L(i)].width = w


def put(ws, row, col, value, font=BLACK, fmt=None, fill=None, bold=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=FONT, color=font.color, bold=bold or font.bold)
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


# ============================================================
# Sheet 1: Assumptions
# ============================================================
A = wb.active
A.title = "Assumptions"
widths(A, [46, 14, 14, 14, 14, 14, 70])
A["A1"] = "Kestrel Industrial Components — TCO Assumptions (ILLUSTRATIVE)"
A["A1"].font = TITLE
A["A2"] = ("Legend: blue = input you can change · black = formula · green = link to another sheet · "
           "yellow fill = quote-based or key assumption to replace with a real quote. "
           "Year 1 = Oct 2026–Sep 2027 (M1–M12).")
A["A2"].font = NOTE

# --- Timeline block (year-varying inputs) ---
A["A4"] = "Rollout ramp and value phasing"
A["A4"].font = SECTION
hdr(A, 5, ["Driver"] + YEARS + ["Basis"])
TL = {}
timeline = [
    ("active_plants", "Plants with edge platform live (plant-years)", [2, 10, 12, 12, 12], NUM,
     "Roadmap: pilot Plant 06 from M5 + OEM wave from M9 ≈ 2 plant-years in Y1; all 12 live by M20 (docs/migration-roadmap.md)"),
    ("downtime_red", "Unplanned-downtime reduction achieved", [0.03, 0.20, 0.30, 0.30, 0.30], PCT,
     "Driver 2 target −30% measured at M24 (end Y2); Y2 is a ramp year. Assumption, not a forecast"),
    ("insurer_flag", "Insurer interim uplift avoided (1 = yes)", [0, 1, 1, 1, 1], NUM,
     "Full renewal July 2027 (M10) removes the +40% interim uplift from Y2 onward (ADR-025 staged acceptance)"),
    ("contain_flag", "Targeted containment available (1 = yes)", [0, 1, 1, 1, 1], NUM,
     "OEM-plant genealogy production-ready M14 (Nov 2027) → benefit from Y2"),
]
r = 6
for key, label, vals, fmt, basis in timeline:
    put(A, r, 1, label)
    for j, v in enumerate(vals):
        put(A, r, 2 + j, v, BLUE, fmt)
    put(A, r, 7, basis).font = NOTE
    TL[key] = r
    r += 1

# --- Single-value inputs ---
r += 1
INP = {}


def section(title):
    global r
    A.cell(row=r, column=1, value=title).font = SECTION
    r += 1
    hdr(A, r, ["Input", "Value", "Unit", "", "", "", "Basis / source"])
    r += 1


def inp(key, label, value, unit, basis, fmt=CUR, quote=False):
    global r
    put(A, r, 1, label)
    put(A, r, 2, value, BLUE, fmt, YELLOW if quote else None)
    put(A, r, 3, unit)
    c = put(A, r, 7, basis)
    c.font = NOTE
    INP[key] = f"Assumptions!$B${r}"
    r += 1


section("Fleet")
inp("plants", "Plants", 12, "plants", "problem-statement.md", NUM)
inp("nodes", "Edge nodes per plant (GDC cluster)", 3, "nodes", "ADR-018: 3-node HA cluster (Azure track also 3, ADR-006)", NUM)
inp("vcpu", "Licensed vCPU per edge node", 16, "vCPU", "gcp-implementation.md edge sizing (16-core servers)", NUM)

section("Edge platform (GCP track) — recurring")
inp("hw_node", "Edge server, per node (capex)", 18000, "$ / node", "Illustrative industrial server 16c/64GB/2×1.92TB NVMe")
inp("gdc_vcpu_mo", "GDC software licence", 24, "$ / vCPU-month", "ILLUSTRATIVE — based on historical GKE Enterprise/Anthos per-vCPU list pricing; replace with Google quote", DEC, True)
inp("mce_n", "Manufacturing Connect edge instances", 16, "instances", "ADR-018: ~10K tags per instance → 2 at larger plants", NUM)
inp("mce_price", "Manufacturing Connect edge licence", 15000, "$ / instance-yr", "ILLUSTRATIVE — sold by Litmus via Marketplace, quote-based", CUR, True)
inp("broker", "Commercial MQTT broker licence", 10000, "$ / plant-yr", "ILLUSTRATIVE — clustered HiveMQ/EMQX-class broker, quote-based", CUR, True)
inp("msp", "Edge managed-service provider (ADR-024 cond. 2)", 4000, "$ / plant-month", "ILLUSTRATIVE — 24×7 edge operations covering GDC + MCe + broker", CUR, True)

section("Security track — recurring (platform-independent)")
inp("sra", "Secure remote access (brokered, MFA)", 8000, "$ / plant-yr", "ILLUSTRATIVE — ADR-003")
inp("otmon", "Passive OT network monitoring", 20000, "$ / plant-yr", "ILLUSTRATIVE — ADR-003")
inp("soc", "Managed OT security monitoring (alert destination)", 150000, "$ / yr", "ILLUSTRATIVE — ADR-003 consequence (no internal OT-security function)")
inp("ramos", "Ramos Arizpe second WAN circuit", 24000, "$ / yr", "ILLUSTRATIVE — ADR-006/ADR-025 WAN resilience")

section("Cloud (GCP) — volume drivers and prices")
inp("vps", "Telemetry values per second (full fleet)", 60000, "values/s", "NFR-3", NUM)
inp("bpv", "Bytes per value on the wire (incl. envelope)", 40, "bytes", "Assumption: batched Sparkplug B/JSON with headers", NUM)
inp("pubsub_mult", "Pub/Sub billable passes (1 publish + subscriptions)", 4, "×", "1 publish + MDE + hot path + archive subscriptions", NUM)
inp("pubsub_tib", "Pub/Sub throughput price", 40, "$ / TiB", "ILLUSTRATIVE list-price basis; confirm in pricing calculator")
inp("df_workers", "Dataflow streaming workers (US + EU)", 11, "workers", "MDE pipeline + hot path, two regions", NUM)
inp("df_hr", "Dataflow cost per worker-hour (incl. streaming engine)", 0.35, "$ / hr", "ILLUSTRATIVE", DEC)
inp("bt_nodes", "Bigtable nodes (US 3 + EU 1)", 4, "nodes", "ADR-021", NUM)
inp("bt_hr", "Bigtable node price", 0.65, "$ / node-hr", "ILLUSTRATIVE list-price basis", DEC)
inp("bt_tb", "Bigtable hot storage (90 days)", 4.7, "TB", "~5.2B values/day × 90 days × ~10 B compressed", DEC)
inp("bt_gb", "Bigtable SSD storage price", 0.17, "$ / GB-month", "ILLUSTRATIVE", DEC)
inp("bq", "BigQuery (storage + queries, both regions)", 30000, "$ / yr", "ILLUSTRATIVE")
inp("csql", "Cloud SQL Enterprise Plus HA (2 regions)", 25000, "$ / yr", "ILLUSTRATIVE — ADR-022")
inp("vertex", "Vertex AI (training + online endpoints)", 40000, "$ / yr", "ILLUSTRATIVE")
inp("looker", "Looker licences", 60000, "$ / yr", "ILLUSTRATIVE — quote-based", CUR, True)
inp("net", "HA VPN tunnels + egress", 10000, "$ / yr", "ILLUSTRATIVE — ADR-023")
inp("obs", "Cloud Logging / Monitoring", 15000, "$ / yr", "ILLUSTRATIVE")
inp("gcs", "Cloud Storage (Bucket Lock archive, evidence project)", 5000, "$ / yr", "ILLUSTRATIVE — ADR-022, ADR-024 cond. 1")

section("Value case")
inp("dt_cost", "Current annual cost of unplanned downtime", 31000000, "$ / yr", "problem-statement.md (Finance estimate)")
inp("premium", "Base cyber-insurance premium", 1800000, "$ / yr", "ILLUSTRATIVE — not stated in problem-statement.md", CUR, True)
inp("uplift", "Interim premium uplift", 0.40, "%", "problem-statement.md (+~40% interim)", PCT)
inp("events", "Field containment events per year", 0.5, "events / yr", "ILLUSTRATIVE — one every two years", DEC)
inp("broad", "Parts contained without serial genealogy", 41000, "parts", "problem-statement.md (2025 containment)", NUM)
inp("targeted", "Parts contained with serial genealogy", 3000, "parts", "problem-statement.md (~3,000 actually suspect)", NUM)
inp("per_part", "Cost per part contained (sort, hold, freight, rework)", 30, "$ / part", "ILLUSTRATIVE", CUR)

section("ADR-024 cost review trigger")
inp("threshold", "Trigger threshold (GCP edge licence stack ÷ Azure)", 1.20, "×", "ADR-024: reopen if GCP stack >20% above Azure over 5 years", DEC)
inp("aio", "Azure IoT Operations fee per node", 600, "$ / node-month", "SCENARIO INPUT — Microsoft's pricing page shows no public price; replace with quote", CUR, True)
inp("gw_n", "Protocol gateways needed on Azure track", 3, "gateways", "Plant 07 + the two Windows 7 HMI plants (MCe replaces them on GCP)", NUM)
inp("gw_price", "Protocol gateway licence", 5000, "$ / gateway-yr", "ILLUSTRATIVE", CUR)


def T(key, j):
    """Timeline cell for year index j (0..4)."""
    return f"Assumptions!${L(2 + j)}${TL[key]}"


def I(key):
    return INP[key]


# ============================================================
# Sheet 2: CloudRunRate (full-fleet annual)
# ============================================================
C = wb.create_sheet("CloudRunRate")
widths(C, [50, 18, 70])
C["A1"] = "GCP Cloud Run-Rate at Full Fleet (annual, illustrative)"
C["A1"].font = TITLE
hdr(C, 3, ["Line item", "Annual $", "Formula basis"])
cloud_var = [
    ("Pub/Sub throughput",
     f"={I('vps')}*{I('bpv')}*86400*365/1099511627776*{I('pubsub_mult')}*{I('pubsub_tib')}",
     "values/s × bytes × seconds/yr ÷ 2^40 × passes × $/TiB"),
    ("Dataflow streaming (MDE + hot path)", f"={I('df_workers')}*{I('df_hr')}*8760", "workers × $/hr × 8,760"),
    ("Bigtable nodes", f"={I('bt_nodes')}*{I('bt_hr')}*8760", "nodes × $/hr × 8,760"),
    ("Bigtable storage", f"={I('bt_tb')}*1000*{I('bt_gb')}*12", "TB × 1,000 × $/GB-month × 12"),
    ("BigQuery", f"={I('bq')}", "input"),
]
cloud_fix = [
    ("Cloud SQL (genealogy)", f"={I('csql')}"),
    ("Vertex AI", f"={I('vertex')}"),
    ("Looker", f"={I('looker')}"),
    ("HA VPN + egress", f"={I('net')}"),
    ("Logging / Monitoring", f"={I('obs')}"),
    ("Cloud Storage archive", f"={I('gcs')}"),
]
row = 4
put(C, row, 1, "Volume-driven (scales with plants live)", bold=True, fill=SUB)
put(C, row, 2, None, fill=SUB); put(C, row, 3, None, fill=SUB)
row += 1
v_start = row
for label, f, basis in cloud_var:
    put(C, row, 1, label); put(C, row, 2, f, BLACK, CUR); put(C, row, 3, basis).font = NOTE
    row += 1
v_end = row - 1
put(C, row, 1, "Subtotal — volume-driven", bold=True, fill=TOTAL)
put(C, row, 2, f"=SUM(B{v_start}:B{v_end})", BLACK, CUR, TOTAL, bold=True)
CLOUD_VAR = f"CloudRunRate!$B${row}"
row += 2
put(C, row, 1, "Fixed baseline (from Y1)", bold=True, fill=SUB)
put(C, row, 2, None, fill=SUB); put(C, row, 3, None, fill=SUB)
row += 1
f_start = row
for label, f in cloud_fix:
    put(C, row, 1, label); put(C, row, 2, f, BLACK, CUR); put(C, row, 3, "input").font = NOTE
    row += 1
f_end = row - 1
put(C, row, 1, "Subtotal — fixed", bold=True, fill=TOTAL)
put(C, row, 2, f"=SUM(B{f_start}:B{f_end})", BLACK, CUR, TOTAL, bold=True)
CLOUD_FIX = f"CloudRunRate!$B${row}"
row += 2
put(C, row, 1, "Total cloud run-rate at full fleet", bold=True, fill=TOTAL)
put(C, row, 2, f"={CLOUD_VAR}+{CLOUD_FIX}", BLACK, CUR, TOTAL, bold=True)

# ============================================================
# Sheet 3: OneTime (program costs with phasing)
# ============================================================
O = wb.create_sheet("OneTime")
widths(O, [44, 12, 12, 14, 16] + [11] * 5 + [14] * 5 + [50])
O["A1"] = "One-Time Program Costs and Phasing (illustrative)"
O["A1"].font = TITLE
O["A2"] = "Blue = inputs. Phasing % per year must sum to 100% (check column Q)."
O["A2"].font = NOTE
hdr(O, 4, ["Item", "Quantity", "Unit cost ($)", "Unit", "Total ($)"] +
    [f"{y} %" for y in YEARS] + [f"{y} $" for y in YEARS] + ["Basis"])
onetime = [
    ("Phase 0 security integrator crews", 12, 120000, "per plant", [1, 0, 0, 0, 0], "ADR-025: 0a/0b/0c zoning, all plants in Y1"),
    ("DMZ firewalls + hosts", 12, 45000, "per plant", [1, 0, 0, 0, 0], "ADR-003 DMZ at every plant"),
    ("GCP landing zone + platform setup (partner)", 1, 250000, "lump sum", [1, 0, 0, 0, 0], "Phase 1"),
    ("Outbox + Kestrel edge services build (4 eng × 9 mo)", 36, 15000, "per eng-month", [1, 0, 0, 0, 0], "ADR-019, ADR-024 cond. 3"),
    ("Vibration/temperature sensors, critical assets", 300, 3500, "per asset", [0.3, 0.7, 0, 0, 0], "~300 critical assets incl. install"),
    ("Tag mapping to ISA-95 model (180K tags × 0.05 h × $120)", 180000, 6, "per tag", [0.3, 0.5, 0.2, 0, 0], "Tiered scope: Tier 1 first (ADR-025)"),
    ("Genealogy station integration, non-MES OEM plants", 3, 150000, "per plant", [1, 0, 0, 0, 0], "Plants 05/08/09, July 2027 shutdown"),
    ("Genealogy parallel run effort", 6, 40000, "per plant", [0.5, 0.5, 0, 0, 0], "ADR-027: 30-day parallel run + 90-day fallback"),
    ("Program management + change", 1, 400000, "lump sum", [0.6, 0.4, 0, 0, 0], "Illustrative"),
]
row = 5
o_start = row
for label, qty, unit_cost, unit, phase, basis in onetime:
    put(O, row, 1, label)
    put(O, row, 2, qty, BLUE, NUM)
    put(O, row, 3, unit_cost, BLUE, CUR)
    put(O, row, 4, unit)
    put(O, row, 5, f"=B{row}*C{row}", BLACK, CUR)
    for j, p in enumerate(phase):
        put(O, row, 6 + j, p, BLUE, PCT)
        put(O, row, 11 + j, f"=$E{row}*{L(6 + j)}{row}", BLACK, CUR)
    put(O, row, 16, basis).font = NOTE
    row += 1
o_end = row - 1
put(O, row, 1, "Total one-time", bold=True, fill=TOTAL)
put(O, row, 5, f"=SUM(E{o_start}:E{o_end})", BLACK, CUR, TOTAL, bold=True)
for j in range(5):
    put(O, row, 11 + j, f"=SUM({L(11 + j)}{o_start}:{L(11 + j)}{o_end})", BLACK, CUR, TOTAL, bold=True)
ONETIME_ROW = row
row += 1
put(O, row, 1, "Phasing check (each row should total 100%)", bold=True)
for rr in range(o_start, o_end + 1):
    put(O, rr, 17, f"=SUM(F{rr}:J{rr})", BLACK, PCT)
O.cell(row=4, column=17, value="Phasing check").font = BOLD_WHITE
O.cell(row=4, column=17).fill = HEADER
O.column_dimensions["Q"].width = 14

# ============================================================
# Sheet 4: FiveYearTCO
# ============================================================
F = wb.create_sheet("FiveYearTCO")
widths(F, [48, 15, 15, 15, 15, 15, 17, 60])
F["A1"] = "Five-Year Cost and Value — GCP Target (illustrative)"
F["A1"].font = TITLE
F["A2"] = "All figures in $. Year 1 = Oct 2026–Sep 2027. Green = links to other sheets."
F["A2"].font = NOTE
hdr(F, 4, ["Line"] + YEARS + ["5-yr total", "Formula basis"])
rows = {}
row = 5


def band(title):
    global row
    for c in range(1, 9):
        put(F, row, c, title if c == 1 else None, bold=True, fill=SUB)
    row += 1


def line(key, label, fn, basis, font=BLACK):
    """fn(j) -> formula string for year j."""
    global row
    put(F, row, 1, label)
    for j in range(5):
        put(F, row, 2 + j, fn(j), font, CUR)
    put(F, row, 7, f"=SUM(B{row}:F{row})", BLACK, CUR)
    put(F, row, 8, basis).font = NOTE
    rows[key] = row
    row += 1


def prev_active(j):
    return "0" if j == 0 else T("active_plants", j - 1)


band("Edge platform (GCP track)")
line("hw", "Edge hardware (capex, as plants go live)",
     lambda j: f"={I('hw_node')}*{I('nodes')}*({T('active_plants', j)}-{prev_active(j)})",
     "$/node × nodes/plant × newly live plants")
line("gdc", "GDC software licence",
     lambda j: f"={I('gdc_vcpu_mo')}*{I('vcpu')}*{I('nodes')}*{T('active_plants', j)}*12",
     "$/vCPU-mo × vCPU × nodes × live plants × 12")
line("mce", "Manufacturing Connect edge licences",
     lambda j: f"={I('mce_n')}*{I('mce_price')}*{T('active_plants', j)}/{I('plants')}",
     "instances × $/yr × share of fleet live")
line("broker", "MQTT broker licences",
     lambda j: f"={I('broker')}*{T('active_plants', j)}", "$/plant-yr × live plants")
line("msp", "Edge managed service (ADR-024 cond. 2)",
     lambda j: f"={I('msp')}*12*{T('active_plants', j)}", "$/plant-mo × 12 × live plants")
band("Security track (platform-independent)")
line("sec", "SRA + OT monitoring + managed OT SOC",
     lambda j: f"=({I('sra')}+{I('otmon')})*{I('plants')}+{I('soc')}",
     "All plants from Y1 (Phase 0a complete by M3)")
line("ramos", "Ramos Arizpe second circuit", lambda j: f"={I('ramos')}", "From Y1")
band("Cloud (GCP)")
line("cvar", "Cloud — volume-driven",
     lambda j: f"={CLOUD_VAR}*{T('active_plants', j)}/{I('plants')}", "Full-fleet run-rate × share live", GREEN)
line("cfix", "Cloud — fixed baseline", lambda j: f"={CLOUD_FIX}", "From Y1", GREEN)
band("One-time program")
line("one", "One-time program costs", lambda j: f"=OneTime!{L(11 + j)}{ONETIME_ROW}", "OneTime sheet phasing", GREEN)

cost_keys = ["hw", "gdc", "mce", "broker", "msp", "sec", "ramos", "cvar", "cfix", "one"]
put(F, row, 1, "TOTAL COST", bold=True, fill=TOTAL)
for j in range(6):
    col = L(2 + j)
    put(F, row, 2 + j, "=" + "+".join(f"{col}{rows[k]}" for k in cost_keys), BLACK, CUR, TOTAL, bold=True)
rows["cost"] = row
row += 2

band("Value (avoided cost)")
line("v_dt", "Unplanned downtime avoided (driver 2)",
     lambda j: f"={I('dt_cost')}*{T('downtime_red', j)}", "$31M × reduction achieved")
line("v_ins", "Insurance interim uplift avoided (driver 1)",
     lambda j: f"={I('premium')}*{I('uplift')}*{T('insurer_flag', j)}", "premium × 40% × flag")
line("v_con", "Containment cost avoided (driver 3)",
     lambda j: f"={I('events')}*({I('broad')}-{I('targeted')})*{I('per_part')}*{T('contain_flag', j)}",
     "events/yr × (41K − 3K parts) × $/part × flag")
put(F, row, 1, "TOTAL VALUE", bold=True, fill=TOTAL)
for j in range(6):
    col = L(2 + j)
    put(F, row, 2 + j, f"={col}{rows['v_dt']}+{col}{rows['v_ins']}+{col}{rows['v_con']}", BLACK, CUR, TOTAL, bold=True)
rows["value"] = row
row += 2

put(F, row, 1, "NET (value − cost)", bold=True, fill=TOTAL)
for j in range(6):
    col = L(2 + j)
    put(F, row, 2 + j, f"={col}{rows['value']}-{col}{rows['cost']}", BLACK, CUR, TOTAL, bold=True)
rows["net"] = row
row += 1
put(F, row, 1, "Cumulative net", bold=True)
for j in range(5):
    col = L(2 + j)
    f = f"={col}{rows['net']}" if j == 0 else f"={L(1 + j)}{row}+{col}{rows['net']}"
    put(F, row, 2 + j, f, BLACK, CUR, bold=True)
rows["cum"] = row
row += 2

put(F, row, 1, "Steady-state recurring run-rate (Y3, excl. one-time and capex)", bold=True)
put(F, row, 4, f"=D{rows['cost']}-D{rows['one']}-D{rows['hw']}", BLACK, CUR, TOTAL, bold=True)
rows["steady"] = row
row += 1
put(F, row, 1, "  of which edge licence stack (GDC + MCe + broker)")
put(F, row, 4, f"=D{rows['gdc']}+D{rows['mce']}+D{rows['broker']}", BLACK, CUR)
row += 1
put(F, row, 1, "  share of steady-state run-rate")
put(F, row, 4, f"=IFERROR(D{row - 1}/D{rows['steady']},0)", BLACK, PCT)
row += 1
put(F, row, 1, "Steady-state value ÷ steady-state run-rate (Y3)")
put(F, row, 4, f"=IFERROR(D{rows['value']}/D{rows['steady']},0)", BLACK, '0.0"x"')
row += 1
put(F, row, 1, "Payback year (first year cumulative net ≥ 0)")
cr = rows["cum"]
put(F, row, 4, f'=IF(B{cr}>=0,"Y1",IF(C{cr}>=0,"Y2",IF(D{cr}>=0,"Y3",IF(E{cr}>=0,"Y4",IF(F{cr}>=0,"Y5","beyond Y5")))))', BLACK)

# chart
ch = BarChart()
ch.type = "col"
ch.title = "Annual Cost vs. Value (illustrative)"
ch.y_axis.title = "$"
ch.x_axis.title = "Year"
data = Reference(F, min_col=1, max_col=6, min_row=rows["cost"], max_row=rows["cost"])
ch.add_data(data, titles_from_data=True, from_rows=True)
data2 = Reference(F, min_col=1, max_col=6, min_row=rows["value"], max_row=rows["value"])
ch.add_data(data2, titles_from_data=True, from_rows=True)
ch.set_categories(Reference(F, min_col=2, max_col=6, min_row=4, max_row=4))
ch.height, ch.width = 8, 16
F.add_chart(ch, f"A{row + 3}")

# ============================================================
# Sheet 5: TriggerTest (ADR-024 cost review trigger)
# ============================================================
G = wb.create_sheet("TriggerTest")
widths(G, [52, 15, 15, 15, 15, 15, 17, 58])
G["A1"] = "ADR-024 Cost Review Trigger — GCP vs. Azure Edge Licence Stack"
G["A1"].font = TITLE
G["A2"] = ("Compares only the edge licence stacks the trigger names. Same node count, ramp and hardware on both "
           "tracks, so those cancel. Azure IoT Operations pricing is not public — the break-even below is the result.")
G["A2"].font = NOTE
hdr(G, 4, ["Line"] + YEARS + ["5-yr total", "Basis"])
g = {}
grow = 5


def gline(key, label, fn, basis, font=BLACK, bold=False, fill=None):
    global grow
    put(G, grow, 1, label, bold=bold, fill=fill)
    for j in range(5):
        put(G, grow, 2 + j, fn(j), font, CUR, fill, bold)
    put(G, grow, 7, f"=SUM(B{grow}:F{grow})", BLACK, CUR, fill, bold)
    put(G, grow, 8, basis).font = NOTE
    g[key] = grow
    grow += 1


gline("gcp", "GCP edge licence stack (GDC + MCe + broker)",
      lambda j: f"=FiveYearTCO!{L(2 + j)}{rows['gdc']}+FiveYearTCO!{L(2 + j)}{rows['mce']}+FiveYearTCO!{L(2 + j)}{rows['broker']}",
      "Links to FiveYearTCO", GREEN, True)
gline("aio", "Azure IoT Operations node fees",
      lambda j: f"={I('aio')}*{I('nodes')}*{T('active_plants', j)}*12", "$/node-mo × nodes × live plants × 12")
gline("gw", "Azure protocol gateways (not needed on GCP)",
      lambda j: f"={I('gw_n')}*{I('gw_price')}*{T('active_plants', j)}/{I('plants')}", "gateways × $/yr × share live")
gline("az", "Azure edge licence stack",
      lambda j: f"={L(2 + j)}{g['aio']}+{L(2 + j)}{g['gw']}", "IoT Operations + gateways", BLACK, True)
grow += 1
put(G, grow, 1, "Ratio: GCP ÷ Azure (5-yr)", bold=True)
put(G, grow, 7, f"=IFERROR(G{g['gcp']}/G{g['az']},0)", BLACK, '0.00"x"', TOTAL, True)
g["ratio"] = grow
grow += 1
put(G, grow, 1, "Trigger threshold", bold=True)
put(G, grow, 7, f"={I('threshold')}", GREEN, '0.00"x"')
grow += 1
put(G, grow, 1, "Does the trigger fire at the scenario Azure price?", bold=True)
put(G, grow, 7, f'=IF(G{g["ratio"]}>{I("threshold")},"YES — reopen ADR-024","No")', BLACK, None, TOTAL, True)
grow += 2
put(G, grow, 1, "Break-even Azure IoT Operations price ($/node-month)", bold=True)
put(G, grow, 7,
    f"=IFERROR((G{g['gcp']}/{I('threshold')}-G{g['gw']})/({I('nodes')}*SUM(Assumptions!$B${TL['active_plants']}:$F${TL['active_plants']})*12),0)",
    BLACK, CUR, YELLOW, True)
g["be"] = grow
put(G, grow, 8, "If Microsoft's quoted IoT Operations price is BELOW this, the GCP stack exceeds Azure's by >20% → trigger fires").font = NOTE
grow += 1
put(G, grow, 1, "Parity Azure price (stacks equal)", bold=True)
put(G, grow, 7,
    f"=IFERROR((G{g['gcp']}-G{g['gw']})/({I('nodes')}*SUM(Assumptions!$B${TL['active_plants']}:$F${TL['active_plants']})*12),0)",
    BLACK, CUR)

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B5" if ws.title in ("FiveYearTCO", "TriggerTest") else None

wb.save("TCO-Analysis.xlsx")
print("saved TCO-Analysis.xlsx")
